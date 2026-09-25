"""
generate_report.py
Exports professional, multi-sheet Excel reports FROM the SQLite database for Asset Tracker.
Includes summary metrics, active loans with overdue flags, full transaction history,
and master asset inventory with custom styling and formatting.
"""

import pandas as pd
from datetime import datetime
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import database as db

REPORTS_DIR = Path(__file__).parent / "reports"


def apply_header_styling(ws, header_color="005C9E", font_color="FFFFFF"):
    """Applies modern corporate styling to Excel header rows."""
    header_fill = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color=font_color)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='medium', color='0B2545')
    )

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    ws.row_dimensions[1].height = 28

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
            # Default cell font and alignment
            if cell.row > 1:
                cell.font = Font(name="Segoe UI", size=10)
                cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)


def generate_full_report() -> Path:
    """Generates a complete multi-sheet Excel report workbook."""
    REPORTS_DIR.mkdir(exist_ok=True)
    db.init_db()
    conn = db.get_connection()

    # 1. Transactions Data
    transactions_df = pd.read_sql_query("""
        SELECT 
            t.id AS "Tx ID",
            t.timestamp AS "Timestamp",
            t.action AS "Action",
            t.laptop_id AS "Laptop ID",
            l.brand AS "Brand",
            l.model AS "Model",
            t.enrollment_no AS "Enrollment No",
            s.name AS "Student Name",
            s.department AS "Department",
            t.expected_return AS "Expected Return",
            t.return_timestamp AS "Returned At",
            t.status AS "Status"
        FROM transactions t
        LEFT JOIN laptops l ON l.laptop_id = t.laptop_id
        LEFT JOIN students s ON s.enrollment_no = t.enrollment_no
        ORDER BY t.id DESC
    """, conn)

    # 2. Currently Borrowed Data
    borrowed_list = db.currently_borrowed_list()
    borrowed_df = pd.DataFrame(borrowed_list)
    if not borrowed_df.empty:
        borrowed_df = borrowed_df[[
            "laptop_id", "brand", "model",
            "enrollment_no", "student_name", "department", "contact",
            "borrowed_at", "duration_str", "is_overdue", "hours_overdue"
        ]].rename(columns={
            "laptop_id": "Laptop ID",
            "brand": "Brand",
            "model": "Model",
            "enrollment_no": "Student ID",
            "student_name": "Student Name",
            "department": "Department",
            "contact": "Contact",
            "borrowed_at": "Borrowed At",
            "duration_str": "Elapsed Time",
            "is_overdue": "Is Overdue",
            "hours_overdue": "Overdue (Hrs)"
        })

    # 3. Master Asset Inventory
    laptops_df = pd.read_sql_query("""
        SELECT 
            laptop_id AS "Laptop ID",
            brand AS "Brand",
            model AS "Model",
            serial_no AS "Serial No",
            status AS "Status",
            created_at AS "Registered At"
        FROM laptops
        ORDER BY laptop_id ASC
    """, conn)

    # 4. Student Directory
    students_df = pd.read_sql_query("""
        SELECT 
            enrollment_no AS "Enrollment No",
            name AS "Name",
            email AS "Email",
            department AS "Department",
            contact AS "Contact",
            status AS "Account Status",
            created_at AS "Registered At"
        FROM students
        ORDER BY enrollment_no ASC
    """, conn)

    summary_kpi = db.get_fleet_summary()
    conn.close()

    timestamp_str = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = REPORTS_DIR / f"Asset_Tracker_Report_{timestamp_str}.xlsx"

    # Write sheets using pandas ExcelWriter
    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        # Summary Overview DataFrame
        summary_rows = [
            {"Metric": "System", "Value": "Asset Tracker"},
            {"Metric": "Institution / Org", "Value": db.get_setting("institution_name", "Asset Tracker System")},
            {"Metric": "Report Generated At", "Value": datetime.now().strftime("%Y-%m-%d %H:%M:%S")},
            {"Metric": "Total Laptop Inventory", "Value": summary_kpi["total_laptops"]},
            {"Metric": "Laptops Currently Available", "Value": summary_kpi["available_laptops"]},
            {"Metric": "Laptops Currently Checked Out", "Value": summary_kpi["borrowed_laptops"]},
            {"Metric": "Laptops Under Maintenance", "Value": summary_kpi["maintenance_laptops"]},
            {"Metric": "Overdue Loans Alert Count", "Value": summary_kpi["overdue_count"]},
            {"Metric": "Total Enrolled Students", "Value": summary_kpi["total_students"]},
            {"Metric": "Today's Transaction Count", "Value": summary_kpi["today_transactions"]},
            {"Metric": "Overdue Policy Limit", "Value": f"{db.get_setting('overdue_hours', '8')} Hours"},
        ]
        summary_df = pd.DataFrame(summary_rows)
        summary_df.to_excel(writer, sheet_name="Executive Summary", index=False)
        
        if not borrowed_df.empty:
            borrowed_df.to_excel(writer, sheet_name="Active Loans (Out)", index=False)
        else:
            pd.DataFrame([{"Notice": "No laptops currently checked out"}]).to_excel(writer, sheet_name="Active Loans (Out)", index=False)

        transactions_df.to_excel(writer, sheet_name="Transaction Logs", index=False)
        laptops_df.to_excel(writer, sheet_name="Laptops Fleet", index=False)
        students_df.to_excel(writer, sheet_name="Student Roster", index=False)

    # Post-process formatting with openpyxl
    wb = openpyxl.load_workbook(filename)
    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        apply_header_styling(ws)

    wb.save(filename)
    return filename


if __name__ == "__main__":
    report_path = generate_full_report()
    print(f"Report successfully created: {report_path}")